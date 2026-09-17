"""Regression tests for the teacher clock-in rollout fixes.

Covers: Perth local dates and times in timesheets, missed clock-outs,
pairing a manual clock-out with no facility, and numeric hours in exports.
"""
import json
from datetime import date, datetime, time, timedelta
from decimal import Decimal
from io import BytesIO
from zoneinfo import ZoneInfo

import openpyxl
from django.contrib.auth import get_user_model
from django.test import Client, TestCase, override_settings
from django.urls import reverse
from django.utils import timezone

from core.models import TeacherAttendance
from core.services.staff_timesheet_service import MAX_SHIFT_HOURS, StaffTimesheetService
from core.views import validate_clock_transition
from facilities.models import Facility

PERTH = ZoneInfo('Australia/Perth')


def perth(day, hour, minute=0):
    return datetime.combine(day, time(hour, minute), tzinfo=PERTH)


@override_settings(SECURE_SSL_REDIRECT=False)
class ClockRolloutFixTests(TestCase):
    def setUp(self):
        user_model = get_user_model()
        self.teacher = user_model.objects.create_user(
            username='rollout-teacher', password='Teacher123!', first_name='Rollout', last_name='Teacher',
            role='teacher', is_active=True, is_active_staff=True,
        )
        self.admin = user_model.objects.create_user(
            username='rollout-admin', password='Admin123!', first_name='Rollout', last_name='Admin',
            role='admin', is_active=True, is_active_staff=True,
        )
        self.facility = Facility.objects.create(
            name='Rollout Studio', address='1 Testing Street', latitude=Decimal('0.0'),
            longitude=Decimal('0.0'), attendance_radius=500, is_active=True,
        )
        self.client = Client()

    def _record(self, clock_type, timestamp, facility='default', source='gps'):
        return TeacherAttendance.objects.create(
            teacher=self.teacher, clock_type=clock_type, source=source, timestamp=timestamp,
            facility=self.facility if facility == 'default' else facility,
            location_verified=source == 'gps',
        )

    def _timesheet(self, start, end):
        return StaffTimesheetService.get_staff_timesheet_data(self.teacher, start, end)

    # --- Perth local time -------------------------------------------------

    def test_timesheet_rows_show_perth_times(self):
        day = date(2026, 9, 8)
        self._record('clock_in', perth(day, 15, 18))
        self._record('clock_out', perth(day, 17, 12))

        rows = self._timesheet(day, day)['paired_records']

        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]['date'], day)
        self.assertEqual(rows[0]['clock_in_time'], time(15, 18))
        self.assertEqual(rows[0]['clock_out_time'], time(17, 12))

    def test_early_morning_clock_in_stays_on_the_perth_date(self):
        # 07:30 in Perth is 23:30 UTC on the previous day
        day = date(2026, 9, 14)
        self._record('clock_in', perth(day, 7, 30))
        self._record('clock_out', perth(day, 11, 0))

        rows = self._timesheet(day, day)['paired_records']

        self.assertEqual(rows[0]['date'], day)
        self.assertEqual(rows[0]['clock_in_time'], time(7, 30))

    def test_is_today_uses_the_perth_date(self):
        local_midnight = timezone.localtime().replace(hour=0, minute=5, second=0, microsecond=0)
        record = self._record('clock_in', local_midnight)
        record.refresh_from_db()  # the database hands back UTC
        self.assertTrue(record.is_today)

    # --- Missed clock-out ---------------------------------------------------

    def test_forgotten_clock_out_is_not_paid_as_a_long_shift(self):
        tuesday, wednesday = date(2026, 9, 15), date(2026, 9, 16)
        self._record('clock_in', perth(tuesday, 9, 0))
        self._record('clock_out', perth(wednesday, 8, 55))

        data = self._timesheet(tuesday, wednesday)

        self.assertEqual(data['summary']['total_hours'], 0)
        self.assertEqual(data['summary']['completed_sessions'], 0)
        self.assertTrue(all(not row['is_complete'] for row in data['paired_records']))

    def test_shift_within_the_limit_still_pairs_across_midnight(self):
        day = date(2026, 9, 18)
        self._record('clock_in', perth(day, 20, 0))
        self._record('clock_out', perth(day + timedelta(days=1), 1, 30))

        data = self._timesheet(day, day + timedelta(days=1))

        self.assertEqual(data['summary']['completed_sessions'], 1)
        self.assertAlmostEqual(data['summary']['total_hours'], 5.5, places=2)

    def test_missed_clock_out_lets_the_teacher_clock_in_for_today(self):
        self._record('clock_in', timezone.now() - timedelta(hours=MAX_SHIFT_HOURS + 1))

        self.assertEqual(validate_clock_transition(self.teacher, 'clock_in'), (True, None))
        allowed, error = validate_clock_transition(self.teacher, 'clock_out')
        self.assertFalse(allowed)
        self.assertIn('tell the office your finish time', error)

    def test_recent_clock_in_still_blocks_a_second_clock_in(self):
        self._record('clock_in', timezone.now() - timedelta(hours=2))

        allowed, error = validate_clock_transition(self.teacher, 'clock_in')

        self.assertFalse(allowed)
        self.assertIn('already clocked in', error)

    def test_clock_page_shows_missed_clock_out_warning_and_clock_in_button(self):
        self._record('clock_in', timezone.now() - timedelta(hours=MAX_SHIFT_HOURS + 3))
        self.client.login(username='rollout-teacher', password='Teacher123!')

        response = self.client.get(reverse('clock'))

        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.context['has_missed_clock_out'])
        self.assertFalse(response.context['is_clocked_in'])
        self.assertContains(response, 'Missed Clock Out')
        self.assertContains(response, 'id="clockInBtn"')
        self.assertNotContains(response, 'id="clockOutBtn"')

    def test_gps_clock_in_succeeds_after_a_missed_clock_out(self):
        self._record('clock_in', timezone.now() - timedelta(hours=MAX_SHIFT_HOURS + 3))
        self.client.login(username='rollout-teacher', password='Teacher123!')

        response = self.client.post(
            reverse('teacher_clock_submit'),
            data=json.dumps({'clock_type': 'clock_in', 'latitude': 0, 'longitude': 0, 'facility_id': self.facility.id}),
            content_type='application/json',
        )

        self.assertEqual(response.status_code, 200, response.content)
        self.assertEqual(TeacherAttendance.objects.filter(teacher=self.teacher, clock_type='clock_in').count(), 2)

    def test_staff_timesheet_labels_missed_and_open_clock_ins_differently(self):
        self._record('clock_in', timezone.now() - timedelta(hours=MAX_SHIFT_HOURS + 20))
        self._record('clock_in', timezone.now() - timedelta(hours=1))
        self.client.login(username='rollout-admin', password='Admin123!')

        start = timezone.localdate() - timedelta(days=3)
        response = self.client.get(
            reverse('accounts:staff_detail', args=[self.teacher.pk]),
            {'timesheet_start': start.isoformat(), 'timesheet_end': timezone.localdate().isoformat()},
        )

        self.assertContains(response, 'Missing clock-out', count=1)
        self.assertContains(response, 'In Progress', count=1)

    def test_edit_form_prefills_values_a_browser_date_input_accepts(self):
        clock_in = self._record('clock_in', perth(date(2026, 9, 16), 16, 2))
        self.client.login(username='rollout-admin', password='Admin123!')
        url = reverse('accounts:staff_attendance_manual_edit', args=[self.teacher.pk, clock_in.pk])

        page = self.client.get(url)
        self.assertContains(page, 'value="2026-09-16"')
        self.assertContains(page, 'value="16:02"')

        # Save the missed clock-out as a full session, sending the pre-filled values back as a browser would
        response = self.client.post(url, {
            'entry_mode': 'full_session',
            'work_date': '2026-09-16',
            'event_time': '16:02',
            'session_start_time': '16:02',
            'session_end_time': '18:30',
            'manual_reason': 'Forgot to clock out.',
        })
        self.assertEqual(response.status_code, 302)
        data = self._timesheet(date(2026, 9, 16), date(2026, 9, 16))
        self.assertEqual(data['summary']['completed_sessions'], 1)
        self.assertAlmostEqual(data['summary']['total_hours'], 2.47, places=2)

    # --- Manual clock-out with no facility ---------------------------------

    def test_manual_clock_out_without_facility_pairs_with_gps_clock_in(self):
        day = date(2026, 9, 15)
        self._record('clock_in', perth(day, 9, 0))
        self._record('clock_out', perth(day, 17, 0), facility=None, source='manual')

        data = self._timesheet(day, day)

        self.assertEqual(data['summary']['completed_sessions'], 1)
        self.assertAlmostEqual(data['summary']['total_hours'], 8.0, places=2)
        self.assertEqual(data['paired_records'][0]['facility'], self.facility)

    # --- Exports --------------------------------------------------------------

    def _load_workbook(self, response):
        self.assertEqual(response.status_code, 200)
        return openpyxl.load_workbook(BytesIO(response.content), data_only=True)

    def test_staff_export_has_perth_times_and_numeric_hours(self):
        day = date(2026, 9, 8)
        self._record('clock_in', perth(day, 15, 18))
        self._record('clock_out', perth(day, 17, 12))
        self.client.login(username='rollout-admin', password='Admin123!')

        response = self.client.get(
            reverse('accounts:staff_timesheet_export', args=[self.teacher.pk]),
            {'start_date': '2026-09-08', 'end_date': '2026-09-08', 'format': 'excel'},
        )
        rows = list(self._load_workbook(response).active.iter_rows(values_only=True))

        header = next(row for row in rows if row and row[0] == 'Date')
        data_row = next(row for row in rows if row and row[0] == '2026-09-08')
        self.assertEqual(data_row[1], '15:18')
        self.assertEqual(data_row[2], '17:12')
        hours = data_row[header.index('Hours')]
        self.assertIsInstance(hours, float)
        self.assertAlmostEqual(hours, 1.9, places=2)
        total_row = next(row for row in rows if row and row[0] == 'Total Hours:')
        self.assertIsInstance(total_row[1], float)

    def test_overview_export_has_numeric_hours(self):
        day = date(2026, 9, 8)
        self._record('clock_in', perth(day, 9, 0))
        self._record('clock_out', perth(day, 12, 30))
        self.client.login(username='rollout-admin', password='Admin123!')

        response = self.client.get(
            reverse('accounts:staff_timesheet_overview_export'),
            {'start_date': '2026-09-08', 'end_date': '2026-09-08', 'format': 'excel'},
        )
        rows = list(self._load_workbook(response).active.iter_rows(values_only=True))

        teacher_row = next(row for row in rows if row and row[0] == 'Rollout Teacher')
        self.assertEqual(teacher_row[1], 3.5)

    # --- Campus defaults ------------------------------------------------------

    def test_new_campus_defaults_to_a_100_metre_radius(self):
        facility = Facility.objects.create(name='New Campus', address='2 Testing Street')
        self.assertEqual(facility.attendance_radius, 100)
